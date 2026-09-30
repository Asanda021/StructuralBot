import os
import tempfile

from app.engine import (
    calculate_slab,
    cut_list_by_diameter,
    estimate_members,
    grid_rebar,
    multi_face_grid_rebar,
    quality_check_members,
    repeated_bar_rebar,
)
from app.exporter import create_excel, create_pdf


def main():
    # Rebar geometry + procurement
    r = repeated_bar_rebar(
        count=10,
        length_each_m=5.0,
        diameter_mm=16,
        bend_m=0.20,
        hook_m=0.15,
        lap_m=0.50,
    )
    assert r["count_bars"] == 10
    assert abs(r["cut_length_each_m"] - 5.85) < 1e-9
    assert len(r["cut_lengths_m"]) == 10
    assert r["weight_kg"] > 0
    assert r["branches"] > 0

    # Bars longer than the 12m stock length are automatically split with
    # a code-based tensile lap; the user does not enter lap/bend/hook lengths.
    long_bar = repeated_bar_rebar(count=1, length_each_m=13.0, diameter_mm=16)
    assert long_bar["splice_count"] == 1
    assert len(long_bar["cut_lengths_m"]) == 2
    assert max(long_bar["cut_lengths_m"]) <= 12.0
    assert long_bar["lap_m"] >= 0.30
    assert long_bar["length_m"] > 13.0

    # Two-direction thermal mesh: one diameter/spacing, bot calculates both directions.
    grid = grid_rebar(8.0, 6.0, 10, 25)
    assert grid["count_bars"] > 0
    assert len(grid["cut_lengths_m"]) == grid["count_bars"]

    # Two-face wall mesh.
    wall = multi_face_grid_rebar(10.0, 3.0, 12, 20, faces=2)
    assert wall["faces"] == 2
    assert wall["length_m"] > 0

    # Typical joist roof: concrete coefficient and derived joists/blocks.
    slab = calculate_slab({
        "length": 10,
        "width": 8,
        "slab_type": "تیرچه تک",
        "joist_spacing_cm": 50,
        "joist_length_m": 10,
        "block_length_m": 1.0,
        "block_width_m": 0.5,
        "thermal_dia": 8,
        "thermal_spacing_cm": 25,
        "openings": [{"length": 1, "width": 2, "count": 1}],
    })
    assert slab["concrete_coeff"] == 0.18
    assert slab["joist_count"] > 0
    assert slab["block_count"] > 0
    assert slab["thermal"]["length_m"] > 0
    assert slab["opening_area_m2"] == 2

    # Foundation mesh: repeated units must preserve explicit cut pieces.
    import bot
    foundation = bot.calc_member("فونداسیون", "پی منفرد",
        [2, 2.0, 1.8, 0.45, 12, 20, 10, 25, 4, 1.0, 16])
    concrete = next(x for x in foundation if x["name"] == "بتن فونداسیون")
    bottom = next(x for x in foundation if x["name"] == "میلگرد شبکه پایین - X - تعداد قطعه")
    top = next(x for x in foundation if x["name"] == "میلگرد شبکه بالا - X - تعداد قطعه")
    assert abs(concrete["value"] - 3.24) < 1e-9
    assert bottom["value"] > 0
    assert top["value"] > 0

    # U-Boot/waffle geometric takeoff path is expected to be supplied by bot.
    # Here we validate the generic slab engine remains usable for solid slabs.
    solid = calculate_slab({
        "length": 10,
        "width": 8,
        "slab_type": "دال بتنی",
        "thickness": 0.15,
        "thermal_dia": 8,
        "thermal_spacing_cm": 25,
    })
    assert abs(solid["concrete_m3"] - 12.0) < 1e-9

    # Cut list must use explicit piece lengths and never invent them.
    members = [{
        "section": "تیر",
        "member": "B1",
        "components": [{
            "name": "میلگرد پایینی - طول اجرا",
            "value": 17.55,
            "unit": "m",
            "category": "میلگرد",
            "diameter_mm": 16,
            "cut_lengths_m": [5.85] * 3,
        }]
    }]
    result = estimate_members(members)
    cl = result["cut_list"]["16.0"]
    assert cl["pieces_count"] == 3
    assert cl["stock_bars"] == 2
    assert abs(cl["used_length_m"] - 17.55) < 1e-9
    assert result["concrete_total_m3"] == 0

    # QA must flag missing concrete/rebar rather than silently inventing values.
    qa = quality_check_members([{
        "section": "ستون",
        "member": "C1",
        "type": "0.35×0.35",
        "components": [],
    }])
    assert not qa["ok"]
    assert qa["warnings"]

    # Stair takeoff must not invent square geometry for reinforcement.
    stair = __import__("bot").calc_member("پله", "پله بتنی", [1, 12.0, 0.18, 12, 20, 8, 25])
    assert any(x["name"] == "میلگرد اصلی - قطر" for x in stair)
    assert not any(x["name"] == "میلگرد اصلی - تعداد قطعه" for x in stair)

    # QA must flag a rebar piece longer than the default 12m stock bar.
    oversize = quality_check_members([{
        "section": "تیر",
        "member": "B2",
        "components": [{
            "name": "میلگرد پایینی - طول اجرا",
            "value": 13.0,
            "unit": "m",
            "category": "میلگرد",
            "diameter_mm": 16,
            "cut_lengths_m": [13.0],
        }, {
            "name": "بتن تیر",
            "value": 1.0,
            "unit": "m³",
            "category": "بتن",
        }],
    }])
    assert not oversize["ok"]
    assert any("۱۲m" in w for w in oversize["warnings"])

    # Exporters must generate valid files from the current result schema.
    full_result = estimate_members([{
        "section": "سقف",
        "member": "S1",
        "type": "تیرچه تک",
        "components": [
            {"name": "بتن سقف", "value": 12.0, "unit": "m³", "category": "بتن"},
            {"name": "میلگرد حرارتی - طول اجرا", "value": 100.0, "unit": "m",
             "category": "میلگرد", "diameter_mm": 8,
             "cut_lengths_m": [10.0] * 10},
            {"name": "میلگرد حرارتی - وزن اجرا", "value": 39.506,
             "unit": "kg", "category": "میلگرد", "diameter_mm": 8},
            {"name": "میلگرد حرارتی - شاخه خرید", "value": 9,
             "unit": "شاخه", "category": "میلگرد", "diameter_mm": 8,
             "procurement_weight_kg": 42.667},
        ],
    }])
    # Export QA sheet must be generated for professional review.
    qa_result = estimate_members([{
        "section": "ستون", "member": "C2",
        "components": [{
            "name": "بتن ستون", "value": 1.0, "unit": "m³", "category": "بتن"
        }]
    }])
    with tempfile.TemporaryDirectory() as d:
        qa_xlsx = create_excel(qa_result, "QA Test", os.path.join(d, "qa.xlsx"))
        from openpyxl import load_workbook
        book = load_workbook(qa_xlsx, read_only=True)
        assert "QA Review" in book.sheetnames
        book.close()

    with tempfile.TemporaryDirectory() as d:
        xlsx = create_excel(full_result, "Smoke Test", os.path.join(d, "smoke.xlsx"))
        pdf = create_pdf(full_result, "Smoke Test", os.path.join(d, "smoke.pdf"))
        assert os.path.getsize(xlsx) > 0
        assert os.path.getsize(pdf) > 0

    print("STRUCTURALBOT_SMOKE_OK")


if __name__ == "__main__":
    main()


# Extended legacy UX smoke checks
from app.db import Database
import tempfile, os
_fd, _p = tempfile.mkstemp(suffix=".db"); os.close(_fd)
try:
    _db=Database(_p); _db.init(); _db.set_settings(1, unit_system="metric", concrete_grade="C30", rebar_grade="A3", standard="iran", stock_length_m=12)
    _s=_db.settings(1)
    assert _s["concrete_grade"]=="C30" and _s["stock_length_m"]==12.0
    print("settings_extended OK")
finally:
    try: os.remove(_p)
    except OSError: pass
