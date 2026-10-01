from collections import defaultdict
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, PageBreak
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfbase import pdfmetrics
from reportlab.lib.enums import TA_CENTER, TA_RIGHT
import os
import re
import arabic_reshaper
from bidi.algorithm import get_display
from app.i18n import L, item_label, lang_code
from app.foundation_takeoff import foundation_bar_marks

def _rtl_pdf_text(text):
    """Shape Arabic/Persian text and apply bidi ordering for ReportLab."""
    s=str(text or "")
    if not re.search(r"[\u0600-\u06FF\u0750-\u077F]", s):
        return s
    try:
        return get_display(arabic_reshaper.reshape(s))
    except Exception:
        return s

def _en_label(text):
    s=str(text or "")
    repl={
        "فونداسیون":"Foundation","ستون":"Column","تیر":"Beam","سقف":"Slab/Roof","دیوار":"Wall","پله":"Stair",
        "آرماتور":"Reinforcement","بتن":"Concrete","میلگرد":"Rebar","شبکه پایین - دو جهت":"Bottom Reinforcement",
        "شبکه بالا - دو جهت":"Top Reinforcement","شبکه حرارتی دو جهت":"Thermal Reinforcement",
        "میلگرد پایین - راستای طول":"Bottom Reinforcement - Longitudinal",
        "میلگرد پایین - راستای عرض":"Bottom Reinforcement - Transverse",
        "میلگرد بالا - راستای طول":"Top Reinforcement - Longitudinal",
        "میلگرد بالا - راستای عرض":"Top Reinforcement - Transverse",
        "میلگرد طولی":"Longitudinal Reinforcement",
        "میلگرد عرضی":"Transverse Reinforcement",
        "میلگرد انتظار ستون":"Column Starter Bars",
        "انتظار راه‌پله":"Stair Starter Bars",
        "چاله آسانسور":"Elevator Pit","میلگرد پایینی تیر":"Beam Bottom Reinforcement",
        "میلگرد بالایی تیر":"Beam Top Reinforcement","خاموت":"Stirrups","سنجاقی ستون":"Column Crossties",
        "سنجاقی تیر":"Beam Crossties","کمرکش تیر":"Beam Side Bars","میلگرد تقویتی":"Additional Reinforcement",
        "میلگرد انتظار":"Starter Bars","کلاف/ژوئن":"Tie / Joint Reinforcement","سنجاقی ژوئن":"Tie / Joint Crossties",
        "میلگرد منفی":"Negative Reinforcement","اتکا/ادکا":"Support Bars","میلگرد قائم دو وجه":"Vertical Wall Reinforcement",
        "میلگرد افقی دو وجه":"Horizontal Wall Reinforcement","خاموت شناژ":"Tie Beam Stirrups",
        "میلگرد انتظار شناژ":"Tie Beam Starters","شناژ":"Tie Beam","بتن مگر":"Lean Concrete",
        "مساحت":"Area","حجم":"Volume","تعداد":"Count","طول":"Length","عرض":"Width","ضخامت":"Thickness"
    }
    for fa,en in sorted(repl.items(),key=lambda x:-len(x[0])):
        s=s.replace(fa,en)
    return s

def rows(result, lang='fa'):
    out=[]
    for r in result.get("items",[]):
        value=r.get("value",r.get("quantity",0))
        name=r.get("name",r.get("member",""))
        out.append([str(r.get("row","")),item_label(r.get("section",""),lang),item_label(r.get("member",""),lang),item_label(name,lang),
                    f"{float(value):,.3f}",r.get("unit",""),""])
    return out

def _rebar_type_rows(result,lang):
    groups=defaultdict(lambda: {"pieces":0,"length":0.0,"weight":0.0,"branches":0,"buy_weight":0.0,"buy_length":0.0})
    for m in result.get("members",[]):
        for c in m.get("components",[]):
            if c.get("category")!="میلگرد": continue
            name=str(c.get("name",""))
            dia=c.get("diameter_mm")
            if dia is None: continue
            base=name.split(" - ")[0]
            g=groups[(base,float(dia))]
            if name.endswith(" - تعداد قطعه"): g["pieces"] += int(c.get("value",0))
            elif name.endswith(" - طول اجرا"):
                g["length"] += float(c.get("value",0))
            elif name.endswith(" - وزن اجرا"):
                g["weight"] += float(c.get("value",0))
            elif name.endswith(" - شاخه خرید"):
                g["branches"] += int(c.get("value",0))
                g["buy_weight"] += float(c.get("procurement_weight_kg",0))
            elif name.endswith(" - طول خرید"):
                g["buy_length"] += float(c.get("value",0))
    return [(k[0],k[1],v) for k,v in groups.items()]

def _foundation_bar_mark_rows(result, lang):
    marks = result.get("foundation_bar_marks")
    if marks is None and result.get("reinforcement"):
        marks = foundation_bar_marks(result)
    if not marks:
        return []
    out = []
    for m in marks:
        cuts = ", ".join(f"{float(x):.2f}" for x in (m.get("cut_lengths_m") or []))
        out.append([
            m.get("bar_mark",""),
            item_label(m.get("name",""),lang),
            f"Φ{float(m.get("diameter_mm",0)):g}",
            f"{float(m.get("spacing_cm",0)):.0f}" if m.get("spacing_cm") is not None else "",
            item_label(m.get("direction",""),lang),
            int(m.get("count",0)),
            cuts,
            f"{float(m.get("length_m",0)):.2f}",
            f"{float(m.get("weight_kg",0)):.2f}",
            int(m.get("branches",0)),
            f"{float(m.get("procurement_length_m",0)):.2f}",
            f"{float(m.get("procurement_weight_kg",0)):.2f}",
        ])
    return out

def _rebar_diameter_rows(result):
    out=[]
    for dia,data in result.get("rebar_by_diameter",{}).items():
        out.append([f"Φ{dia}",data.get("length_m",0),data.get("weight_kg",0),
                    data.get("branches",0),data.get("procurement_length_m",0),
                    data.get("procurement_weight_kg",0)])
    return out

def _style_sheet(ws, lang="fa"):
    rtl = lang_code(lang) in ("fa","ar")
    ws.sheet_view.rightToLeft = rtl
    thin=Side(style="thin",color="B7B7B7")
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment=Alignment(vertical="top",wrap_text=True,horizontal="right" if rtl else "left")
            cell.border=Border(bottom=thin)
    for cell in ws[1]:
        cell.font=Font(bold=True)
    ws.freeze_panes="A2"
    ws.auto_filter.ref=ws.dimensions
    for col in ws.columns:
        letter=get_column_letter(col[0].column)
        width=min(max(max(len(str(c.value or "")) for c in col)+2,12),42)
        ws.column_dimensions[letter].width=width

def create_excel(result,project_name,path,lang=None):
    lang=lang_code(lang or result.get('project_settings',{}).get('language','fa'))
    wb=Workbook()
    ws=wb.active; ws.title=L("summary",lang)[:31]
    ws.append([L("report",lang)]); ws.append([L("project",lang),project_name])
    ws.append([L("members",lang),result.get("member_count",0)])
    ws.append([L("concrete",lang),result.get("concrete_total_m3",0)])
    total_rebar=sum(float(x.get("weight_kg",0)) for x in result.get("rebar_by_diameter",{}).values())
    total_buy=sum(float(x.get("procurement_weight_kg",0)) for x in result.get("rebar_by_diameter",{}).values())
    total_bars=sum(int(x.get("branches",0)) for x in result.get("rebar_by_diameter",{}).values())
    ws.append([L("rebar_exec",lang),total_rebar])
    ws.append([L("rebar_buy",lang),total_buy])
    ws.append([L("stock",lang),total_bars])
    ps=result.get("project_settings",{})
    ws.append([L("standard",lang),ps.get("standard","")])
    ws.append([L("language",lang),ps.get("language","")])
    _style_sheet(ws,lang)

    detail=wb.create_sheet(L("detail",lang)[:31])
    detail.append([L("row_no",lang),L("section",lang),L("member",lang),L("item",lang),L("quantity",lang),L("unit",lang),L("notes",lang)])
    for row in rows(result,lang): detail.append(row)
    _style_sheet(detail,lang)

    bm_rows = _foundation_bar_mark_rows(result, lang)
    if bm_rows:
        bm=wb.create_sheet(L("bar_mark",lang)[:31])
        bm.append([L("bar_mark",lang),L("item",lang),L("dia",lang),L("spacing",lang),L("direction",lang),
                   L("pieces",lang),L("cut_lengths",lang),L("exec_len",lang),L("exec_weight",lang),
                   L("stock_bars",lang),L("buy_len",lang),L("buy_weight",lang)])
        for row in bm_rows: bm.append(row)
        _style_sheet(bm,lang)

    rb=wb.create_sheet(L("rebar_type",lang)[:31])
    rb.append([L("rebar_type",lang),L("dia",lang),L("pieces",lang),L("exec_len",lang),L("exec_weight",lang),L("stock_bars",lang),L("buy_len",lang),L("buy_weight",lang)])
    for base,dia,g in _rebar_type_rows(result,lang):
        rb.append([item_label(base,lang),f"Φ{dia:g}",g["pieces"],g["length"],g["weight"],g["branches"],g["buy_length"],g["buy_weight"]])
    _style_sheet(rb,lang)

    rd=wb.create_sheet(L("procurement",lang)[:31])
    rd.append([L("dia",lang),L("exec_len",lang),L("exec_weight",lang),L("stock_bars",lang),L("buy_len",lang),L("buy_weight",lang)])
    for row in _rebar_diameter_rows(result): rd.append(row)
    _style_sheet(rd,lang)

    concrete=wb.create_sheet(L("concrete",lang)[:31])
    concrete.append([L("project",lang),project_name])
    concrete.append([L("concrete",lang),result.get("concrete_total_m3",0)])
    concrete.append([L("members",lang),result.get("member_count",0)])
    _style_sheet(concrete,lang)

    qa_ws=wb.create_sheet(L("qa",lang)[:31])
    qa_ws.append([L("status",lang),L("notes",lang)])
    qa=result.get("qa",{})
    if qa.get("ok"): qa_ws.append([L("ok",lang),L("ok",lang)])
    else:
        for warning in qa.get("warnings",[]): qa_ws.append([L("warning",lang),item_label(warning,lang)])
    _style_sheet(qa_ws,lang)

    cl=wb.create_sheet(L("cut",lang)[:31])
    cl.append([L("dia",lang),L("stock_length",lang),L("pieces",lang),L("used",lang),L("waste",lang)])
    for dia,data in result.get("cut_list",{}).items():
        cl.append([f"Φ{dia}",data.get("stock_bars",0),data.get("pieces_count",0),
                   data.get("used_length_m",0),data.get("waste_length_m",0)])
    _style_sheet(cl,lang)

    units=wb.create_sheet(L("unit_totals",lang)[:31])
    units.append([L("unit",lang),L("quantity",lang)])
    for k,v in result.get("totals_by_unit",{}).items(): units.append([k,v])
    _style_sheet(units,lang)

    ai= result.get("ai_explanation")
    if ai:
        aiws=wb.create_sheet(L("ai",lang)[:31])
        aiws.append([L("ai",lang)]); aiws.append([ai])
        aiws.column_dimensions["A"].width=110
        aiws["A2"].alignment=Alignment(wrap_text=True,vertical="top")

    wb.save(path); return path

def _pdf_font(lang="fa"):
    if lang_code(lang)=="zh":
        try:
            from reportlab.pdfbase.cidfonts import UnicodeCIDFont
            pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
            return "STSong-Light"
        except Exception: pass
    candidates=[
        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
        "/usr/share/fonts/truetype/dejavu/DejaVuSansCondensed.ttf",
    ]
    for p in candidates:
        if os.path.exists(p):
            try:
                pdfmetrics.registerFont(TTFont("StructuralBotFont",p))
                return "StructuralBotFont"
            except Exception: pass
    return "Helvetica"

def create_pdf(result,project_name,path,lang=None):
    lang=lang_code(lang or result.get('project_settings',{}).get('language','fa'))
    doc=SimpleDocTemplate(path,pagesize=landscape(A4),rightMargin=24,leftMargin=24,topMargin=24,bottomMargin=24)
    font=_pdf_font(lang)
    styles=getSampleStyleSheet()
    title=ParagraphStyle("SBTitle",parent=styles["Title"],fontName=font,alignment=TA_CENTER,fontSize=18,leading=22)
    body=ParagraphStyle("SBBody",parent=styles["BodyText"],fontName=font,fontSize=8.5,leading=11)
    head=ParagraphStyle("SBHead",parent=styles["Heading2"],fontName=font,fontSize=13,leading=16,alignment=TA_RIGHT)
    story=[Paragraph(_rtl_pdf_text(L("report",lang)),title),
           Paragraph(_rtl_pdf_text(f"{L('project',lang)}: {project_name}"),body),Spacer(1,10)]

    total_rebar=sum(float(x.get("weight_kg",0)) for x in result.get("rebar_by_diameter",{}).values())
    total_buy=sum(float(x.get("procurement_weight_kg",0)) for x in result.get("rebar_by_diameter",{}).values())
    total_bars=sum(int(x.get("branches",0)) for x in result.get("rebar_by_diameter",{}).values())
    summary=[
        [L("members",lang),L("concrete",lang),L("rebar_exec",lang),L("stock",lang),L("rebar_buy",lang)],
        [result.get("member_count",0),f"{result.get('concrete_total_m3',0):,.3f}",f"{total_rebar:,.2f}",total_bars,f"{total_buy:,.2f}"]
    ]
    t=Table(summary,colWidths=[70,95,120,80,120])
    t.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,0),colors.lightgrey),("GRID",(0,0),(-1,-1),0.5,colors.grey),("FONTNAME",(0,0),(-1,-1),font),("ALIGN",(0,0),(-1,-1),"CENTER")]))
    story += [t,Spacer(1,14),Paragraph(_rtl_pdf_text(L("detail",lang)),head)]
    data=[[L("row_no",lang),L("section",lang),L("member",lang),L("item",lang),L("quantity",lang),L("unit",lang),L("notes",lang)]]+rows(result,lang)
    data=[[Paragraph(_rtl_pdf_text(x),body) for x in row] for row in data]
    table=Table(data,repeatRows=1,colWidths=[30,60,75,180,65,45,180])
    table.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,0),colors.lightgrey),("GRID",(0,0),(-1,-1),0.35,colors.grey),("FONTNAME",(0,0),(-1,-1),font),("VALIGN",(0,0),(-1,-1),"TOP")]))
    story += [table,PageBreak(),Paragraph(_rtl_pdf_text(L("rebar_type",lang)),head)]
    rb=[[L("rebar_type",lang),L("dia",lang),L("pieces",lang),L("exec_len",lang),L("exec_weight",lang),L("stock_bars",lang),L("buy_len",lang),L("buy_weight",lang)]]
    for base,dia,g in _rebar_type_rows(result,lang):
        rb.append([item_label(base,lang),f"Φ{dia:g}",g["pieces"],f"{g['length']:.2f}",f"{g['weight']:.2f}",g["branches"],f"{g['buy_length']:.2f}",f"{g['buy_weight']:.2f}"])
    rb=[[Paragraph(_rtl_pdf_text(x),body) for x in row] for row in rb]
    rt=Table(rb,repeatRows=1,colWidths=[150,45,65,75,75,65,75,75])
    rt.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,0),colors.lightgrey),("GRID",(0,0),(-1,-1),0.35,colors.grey),("FONTNAME",(0,0),(-1,-1),font),("ALIGN",(1,1),(-1,-1),"CENTER")]))
    bm_rows = _foundation_bar_mark_rows(result, lang)
    if bm_rows:
        story += [rt,PageBreak(),Paragraph(_rtl_pdf_text(L("bar_mark",lang)),head)]
        bm=[[L("bar_mark",lang),L("item",lang),L("dia",lang),L("spacing",lang),L("direction",lang),
             L("pieces",lang),L("cut_lengths",lang),L("exec_len",lang),L("exec_weight",lang),
             L("stock_bars",lang),L("buy_len",lang),L("buy_weight",lang)]]
        bm += bm_rows
        bm=[[Paragraph(_rtl_pdf_text(x),body) for x in row] for row in bm]
        bmt=Table(bm,repeatRows=1,colWidths=[42,112,38,48,50,45,112,60,62,52,60,62])
        bmt.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,0),colors.lightgrey),("GRID",(0,0),(-1,-1),0.35,colors.grey),
                                 ("FONTNAME",(0,0),(-1,-1),font),("VALIGN",(0,0),(-1,-1),"TOP"),
                                 ("ALIGN",(2,1),(-1,-1),"CENTER")]))
        story.append(bmt)
    story += [Spacer(1,14),Paragraph(_rtl_pdf_text(L("procurement",lang)),head)]
    rd=[[L("dia",lang),L("exec_len",lang),L("exec_weight",lang),L("stock_bars",lang),L("buy_len",lang),L("buy_weight",lang)]]
    for row in _rebar_diameter_rows(result): rd.append([str(x) for x in row])
    rd=[[Paragraph(_rtl_pdf_text(x),body) for x in row] for row in rd]
    rdt=Table(rd,repeatRows=1)
    rdt.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,0),colors.lightgrey),("GRID",(0,0),(-1,-1),0.35,colors.grey),("FONTNAME",(0,0),(-1,-1),font),("ALIGN",(0,0),(-1,-1),"CENTER")]))
    story += [rdt,Spacer(1,14),Paragraph(_rtl_pdf_text(L("qa",lang)),head)]
    qa=result.get("qa",{})
    story.append(Paragraph(f"{L('status',lang)}: "+(L('ok',lang) if qa.get("ok") else f"{L('check',lang)} ({len(qa.get('warnings',[]))})"),body))
    for _w in qa.get("warnings",[])[:12]: story.append(Paragraph(_rtl_pdf_text(f"⚠ {L('warning',lang)}: {item_label(_w,lang)}"),body))
    story.append(Spacer(1,8))
    if result.get("ai_explanation"):
        story += [Paragraph(_rtl_pdf_text(L("ai",lang)),head),Paragraph(_rtl_pdf_text(result["ai_explanation"]),body),Spacer(1,8)]
    story.append(Paragraph(_rtl_pdf_text(L("final_note",lang)),body))
    doc.build(story); return path
